from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from database import init_db
from routers.auth import router as auth_router
from routers.saved_jobs import router as saved_jobs_router
from services.job_matcher import match_job
from services.job_service import JobServiceError, search_jobs
from services.resume_analyzer import analyze_resume

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

@asynccontextmanager
async def lifespan(_: FastAPI):
    init_db()
    yield


app = FastAPI(
    lifespan=lifespan,
    title="JayPilot AI API",
    description="Backend API for the JayPilot AI job discovery platform",
    version="1.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["*"],
)

app.include_router(auth_router)
app.include_router(saved_jobs_router)


class CandidateProfileIn(BaseModel):
    skills: list[str] = Field(default_factory=list, max_length=100)
    education: list[dict[str, str]] = Field(default_factory=list, max_length=20)
    experience: list[dict[str, str]] = Field(default_factory=list, max_length=20)
    certifications: list[str] = Field(default_factory=list, max_length=30)
    job_search_terms: list[str] = Field(default_factory=list, max_length=20)


class JobMatchRequest(BaseModel):
    query: str = Field(min_length=1, max_length=100)
    location: str = Field("", max_length=100)
    page: int = Field(1, ge=1, le=50)
    results_per_page: int = Field(10, ge=1, le=50)
    candidate_profile: CandidateProfileIn


def extract_text_from_pdf(content: bytes) -> str:
    with pymupdf.open(stream=content, filetype="pdf") as pdf:
        return "\n".join(page.get_text() for page in pdf).strip()


def extract_text_from_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(part for part in parts if part.strip()).strip()


@app.get("/")
def read_root() -> dict[str, str]:
    return {
        "message": "JayPilot AI API is running",
        "status": "success",
    }


@app.get("/health")
def health_check() -> dict[str, str]:
    return {"status": "healthy"}


@app.get("/jobs/search")
async def jobs_search(
    query: str = Query(..., min_length=1, max_length=100),
    location: str = Query("", max_length=100),
    page: int = Query(1, ge=1, le=50),
    results_per_page: int = Query(10, ge=1, le=50),
) -> dict:
    try:
        return await search_jobs(query.strip(), location.strip(), page, results_per_page)
    except JobServiceError as error:
        raise HTTPException(status_code=error.status_code, detail=error.message)


@app.post("/jobs/match")
async def jobs_match(request: JobMatchRequest) -> dict:
    try:
        result = await search_jobs(
            request.query.strip(),
            request.location.strip(),
            request.page,
            request.results_per_page,
            keep_full_description=True,
        )
    except JobServiceError as error:
        raise HTTPException(status_code=error.status_code, detail=error.message)

    profile = request.candidate_profile.model_dump()
    for job in result["jobs"]:
        full_text = job.pop("full_description")
        job["match"] = match_job(profile, job, full_text)
    result["jobs"].sort(key=lambda j: j["match"]["match_score"], reverse=True)
    return result


@app.post("/resume/upload")
async def upload_resume(file: UploadFile = File(...)) -> dict:
    filename = file.filename or ""
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Unsupported file type. Only PDF and DOCX files are allowed.",
        )

    content = await file.read(MAX_FILE_SIZE + 1)
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413, detail="File is too large. Maximum size is 5 MB."
        )

    try:
        if extension == ".pdf":
            text = extract_text_from_pdf(content)
        else:
            text = extract_text_from_docx(content)
    except Exception:
        raise HTTPException(
            status_code=422,
            detail="Could not read the file. It may be corrupted or password protected.",
        )

    if not text:
        raise HTTPException(
            status_code=422,
            detail="No text could be extracted. The resume may be empty or a scanned image.",
        )

    return {
        "filename": filename,
        "file_type": extension.lstrip("."),
        "extracted_text": text,
        "character_count": len(text),
        "status": "success",
        "candidate_profile": analyze_resume(text),
    }
